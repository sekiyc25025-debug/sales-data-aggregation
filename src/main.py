import csv
import re
import sys
from copy import copy
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.utils.datetime import to_excel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties


print("売上データ自動集計ツール")
print("------------------------")
print("処理を開始します。")


# ============================================================
# パス設定
# ============================================================

if getattr(sys, "frozen", False):
    # exeとして実行している場合
    BASE_DIR = Path(sys.executable).resolve().parent
else:
    # Pythonスクリプトとして実行している場合
    BASE_DIR = Path(__file__).resolve().parent.parent

INPUT_DIR = BASE_DIR / "input"
OUTPUT_DIR = BASE_DIR / "output"

report_path = INPUT_DIR / "製品別売上レポート.xlsx"
output_path = OUTPUT_DIR / "製品別売上レポート.xlsx"


# ============================================================
# エラー終了
# ============================================================

def stop_with_errors(errors):
    """エラー内容を表示し、Enter入力後に処理を終了する。"""
    print()
    print("========================================")
    print("【入力エラー】")
    print("以下の問題を確認してください。")
    print()

    for error in errors:
        print(f"・{error}")

    print()
    print("処理を中止しました。")
    print("問題を修正してから、もう一度実行してください。")
    print("========================================")
    print()

    input("Enterキーを押すと終了します。")
    raise SystemExit(1)


# ============================================================
# 入力ファイル探索・振り分け
# ============================================================

def find_order_files(input_dir):
    """
    inputフォルダから、発注書A / 発注書Bとして扱う
    ファイルを自動検出する。

    A形式：ファイル名が「発注書A」で始まる
           かつ拡張子が.xlsx
    B形式：ファイル名が「発注書B」で始まる
           かつ拡張子が.csv

    拡張子違いは「無視」せず、エラーとして扱う。
    """

    a_files = []
    b_files = []
    errors = []

    for file in sorted(input_dir.iterdir()):
        if not file.is_file():
            continue

        if file.name.startswith("発注書A"):
            if file.suffix.lower() != ".xlsx":
                errors.append(
                    f"{file.name}："
                    "発注書Aは.xlsx形式である必要があります。"
                )
            else:
                a_files.append(file)

        elif file.name.startswith("発注書B"):
            if file.suffix.lower() != ".csv":
                errors.append(
                    f"{file.name}："
                    "発注書Bは.csv形式である必要があります。"
                )
            else:
                b_files.append(file)

    return a_files, b_files, errors


# ============================================================
# 空行チェック
# ============================================================

def is_blank_value(value):
    """セル/CSV項目が空かどうかを判定する。"""
    if value is None:
        return True

    if isinstance(value, str) and value.strip() == "":
        return True

    return False


def find_blank_row_in_xlsx(file):
    """
    Excelファイルについて、ヘッダー以降の
    「データ途中の完全な空行」を検出する。

    先頭・末尾の空行は問題とせず、
    最初と最後のデータ行の間だけ確認する。
    """

    wb = load_workbook(file, read_only=True, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(min_row=2, values_only=True))

    non_blank_row_numbers = []

    for row_number, row in enumerate(rows, start=2):
        if any(not is_blank_value(value) for value in row):
            non_blank_row_numbers.append(row_number)

    if not non_blank_row_numbers:
        return None

    first_data_row = min(non_blank_row_numbers)
    last_data_row = max(non_blank_row_numbers)

    for row_number in range(first_data_row, last_data_row + 1):
        row = rows[row_number - 2]

        if all(is_blank_value(value) for value in row):
            return row_number

    return None


def find_blank_row_in_csv(file):
    """
    CSVファイルについて、ヘッダー以降の
    「データ途中の完全な空行」を検出する。
    """

    with file.open("r", newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))

    non_blank_row_numbers = []

    for row_number, row in enumerate(rows[1:], start=2):
        if any(not is_blank_value(value) for value in row):
            non_blank_row_numbers.append(row_number)

    if not non_blank_row_numbers:
        return None

    first_data_row = min(non_blank_row_numbers)
    last_data_row = max(non_blank_row_numbers)

    for row_number in range(first_data_row, last_data_row + 1):
        row = rows[row_number - 1]

        if all(is_blank_value(value) for value in row):
            return row_number

    return None


# ============================================================
# データ検証
# ============================================================

def validate_common_data(df, file, file_type):
    """
    共通化した列名のDataFrameを検証する。

    Ver.3では、問題を自動修正するのではなく、
    人間が修正できるようエラーを返す。
    """

    errors = []

    if df.empty:
        errors.append(
            f"{file.name}：データが空です。"
            "データ欠損の可能性があるため、内容を確認してください。"
        )
        return errors

    # 必須項目の空欄チェック
    required_columns = [
        "order_date",
        "product_name",
        "quantity",
    ]

    for column in required_columns:
        blank_mask = df[column].isna()

        if df[column].dtype == object:
            blank_mask = blank_mask | (
                df[column].astype(str).str.strip() == ""
            )

        if blank_mask.any():
            row_numbers = (df.index[blank_mask] + 2).tolist()

            errors.append(
                f"{file.name}："
                f"「{column}」に空欄があります。"
                f"（Excel/CSV上の行番号: {row_numbers}）"
            )

    # 数量の数値チェック
    raw_quantity = df["quantity"]

    quantity_numeric = pd.to_numeric(
        raw_quantity,
        errors="coerce"
    )

    invalid_quantity_mask = (
        raw_quantity.isna()
        | (
            raw_quantity.astype(str).str.strip() == ""
        )
        | quantity_numeric.isna()
    )

    if invalid_quantity_mask.any():
        row_numbers = (
            df.index[invalid_quantity_mask] + 2
        ).tolist()

        errors.append(
            f"{file.name}："
            "「quantity」に数値として扱えない値があります。"
            f"（Excel/CSV上の行番号: {row_numbers}）"
        )
    else:
        df["quantity"] = quantity_numeric

    # 日付チェック
    if df["order_date"].isna().any():
        row_numbers = (
            df.index[df["order_date"].isna()] + 2
        ).tolist()

        errors.append(
            f"{file.name}："
            "「order_date」を日付として解釈できない値があります。"
            f"（Excel/CSV上の行番号: {row_numbers}）"
        )

    # 商品名チェックは必須項目チェックで実施済み
    return errors


# ============================================================
# 令和日付変換
# ============================================================

def convert_reiwa_date(value):
    """B形式の令和日付をdatetimeへ変換する。"""

    if pd.isna(value):
        return pd.NaT

    if isinstance(value, (pd.Timestamp, datetime, date)):
        return pd.Timestamp(value)

    text = str(value).strip()

    match = re.fullmatch(
        r"令和(\d+)年(\d+)月(\d+)日",
        text
    )

    if not match:
        return pd.NaT

    reiwa_year = int(match.group(1))
    month = int(match.group(2))
    day = int(match.group(3))

    year = reiwa_year + 2018

    try:
        return pd.Timestamp(
            year=year,
            month=month,
            day=day
        )
    except ValueError:
        return pd.NaT


# ============================================================
# A形式・B形式の読み込み
# ============================================================

def read_a_file(file):
    """A形式のExcelを読み込んで共通列名へ変換する。"""

    blank_row = find_blank_row_in_xlsx(file)

    if blank_row is not None:
        raise ValueError(
            f"データ途中に空行があります。"
            f"（Excel上の行番号: {blank_row}）"
        )

    df = pd.read_excel(file)

    required_columns = [
        "発注日",
        "商品名",
        "数量",
    ]

    missing_columns = [
        column
        for column in required_columns
        if column not in df.columns
    ]

    if missing_columns:
        raise ValueError(
            "A形式に必要な列がありません："
            + "、".join(missing_columns)
            + "。"
            "ファイル内容がA形式と一致しているか確認してください。"
        )

    df = df.rename(columns={
        "発注日": "order_date",
        "商品名": "product_name",
        "数量": "quantity"
    })

    # A形式はExcel上の日付として読み込めない
    # 数値データをそのまま日付と解釈せず、エラーにする。
    numeric_date_mask = pd.api.types.is_numeric_dtype(
        df["order_date"]
    )

    if numeric_date_mask:
        raise ValueError(
            "「発注日」が数値形式で読み込まれています。"
            "日付として保存されているか確認してください。"
        )

    df["order_date"] = pd.to_datetime(
        df["order_date"],
        errors="coerce"
    )

    errors = validate_common_data(
        df,
        file,
        "A"
    )

    if errors:
        raise ValueError("\n".join(errors))

    return df


def read_b_file(file):
    """B形式のCSVを読み込んで共通列名へ変換する。"""

    blank_row = find_blank_row_in_csv(file)

    if blank_row is not None:
        raise ValueError(
            f"データ途中に空行があります。"
            f"（CSV上の行番号: {blank_row}）"
        )

    df = pd.read_csv(
        file,
        encoding="utf-8-sig"
    )

    required_columns = [
        "発注日付",
        "個数",
        "製品名",
    ]

    missing_columns = [
        column
        for column in required_columns
        if column not in df.columns
    ]

    if missing_columns:
        raise ValueError(
            "B形式に必要な列がありません："
            + "、".join(missing_columns)
            + "。"
            "ファイル内容がB形式と一致しているか確認してください。"
        )

    df = df.rename(columns={
        "発注日付": "order_date",
        "製品名": "product_name",
        "個数": "quantity"
    })

    df["order_date"] = df["order_date"].apply(
        convert_reiwa_date
    )

    errors = validate_common_data(
        df,
        file,
        "B"
    )

    if errors:
        raise ValueError("\n".join(errors))

    return df


# ============================================================
# ファイル読み込み
# ============================================================

a_files, b_files, discovery_errors = find_order_files(INPUT_DIR)

order_file_count = len(a_files) + len(b_files)

if order_file_count == 0:
    stop_with_errors([
        "発注書が1件も見つかりません。"
        "inputフォルダを確認してください。"
    ])

print(
    f"発注書候補を {order_file_count} 件検出しました。"
)

print("読み込み対象：")

if a_files:
    print("  A形式")
    for file in a_files:
        print(f"    ・{file.name}")

if b_files:
    print("  B形式")
    for file in b_files:
        print(f"    ・{file.name}")

file_errors = list(discovery_errors)

a_dfs = []
b_dfs = []

print()
print("発注書を読み込んでいます...")

for file in a_files:
    try:
        a_df = read_a_file(file)
        a_dfs.append(a_df)
        print(f"  ✓ {file.name}：A形式")

    except Exception as e:
        file_errors.append(
            f"{file.name}：{e}"
        )

for file in b_files:
    try:
        b_df = read_b_file(file)
        b_dfs.append(b_df)
        print(f"  ✓ {file.name}：B形式")

    except Exception as e:
        file_errors.append(
            f"{file.name}：{e}"
        )

if file_errors:
    stop_with_errors(file_errors)


# ============================================================
# 集計
# ============================================================

print()
print("データを集計しています...")

df_all = pd.concat(
    a_dfs + b_dfs,
    ignore_index=True
)

df_summary = (
    df_all
    .groupby(
        ["order_date", "product_name"],
        as_index=False
    )["quantity"]
    .sum()
)


# ============================================================
# レポート作成
# ============================================================

print("レポートを作成しています...")

OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True
)

if not report_path.exists():
    stop_with_errors([
        f"集計テンプレートが見つかりません："
        f"{report_path.name}"
    ])

wb = load_workbook(report_path)
ws_report = wb.active

start_row = 2
template_end_row = 21

# 通常行と最終行の書式を保存
normal_styles = []
last_styles = []

for col in range(1, 4):
    normal_cell = ws_report.cell(
        row=2,
        column=col
    )

    last_cell = ws_report.cell(
        row=21,
        column=col
    )

    normal_styles.append({
        "font": copy(normal_cell.font),
        "fill": copy(normal_cell.fill),
        "border": copy(normal_cell.border),
        "alignment": copy(normal_cell.alignment),
        "number_format": normal_cell.number_format,
    })

    last_styles.append({
        "font": copy(last_cell.font),
        "fill": copy(last_cell.fill),
        "border": copy(last_cell.border),
        "alignment": copy(last_cell.alignment),
        "number_format": last_cell.number_format,
    })

last_data_row = start_row + len(df_summary) - 1

for row_index, row in enumerate(
    df_summary.itertuples(index=False),
    start=start_row
):
    values = [
        row.order_date,
        row.product_name,
        row.quantity,
    ]

    for col_index, value in enumerate(
        values,
        start=1
    ):
        cell = ws_report.cell(
            row=row_index,
            column=col_index,
            value=value
        )

        style = normal_styles[
            col_index - 1
        ]

        cell.font = copy(style["font"])
        cell.fill = copy(style["fill"])
        cell.border = copy(style["border"])
        cell.alignment = copy(style["alignment"])
        cell.number_format = style["number_format"]

    ws_report.cell(
        row=row_index,
        column=1
    ).number_format = "yyyy/mm/dd"

for col_index in range(1, 4):
    cell = ws_report.cell(
        row=last_data_row,
        column=col_index
    )

    style = last_styles[
        col_index - 1
    ]

    cell.border = copy(style["border"])


# ============================================================
# グラフ作成
# ============================================================

df_chart = (
    df_summary
    .pivot(
        index="order_date",
        columns="product_name",
        values="quantity"
    )
    .fillna(0)
)

if "グラフ用データ" in wb.sheetnames:
    del wb["グラフ用データ"]

ws_data = wb.create_sheet(
    "グラフ用データ"
)

ws_data.cell(
    row=1,
    column=1,
    value="年月日"
)

for col_index, product_name in enumerate(
    df_chart.columns,
    start=2
):
    ws_data.cell(
        row=1,
        column=col_index,
        value=product_name
    )

for row_index, (
    order_date,
    row
) in enumerate(
    df_chart.iterrows(),
    start=2
):
    ws_data.cell(
        row=row_index,
        column=1,
        value=order_date
    )

    ws_data.cell(
        row=row_index,
        column=1
    ).number_format = "yyyy/mm/dd"

    for col_index, value in enumerate(
        row,
        start=2
    ):
        ws_data.cell(
            row=row_index,
            column=col_index,
            value=value
        )

chart = ScatterChart()

chart.title = "商品別 日別売上数量"

chart.x_axis.title = "日付"
chart.y_axis.title = "売上数量"

chart.x_axis.delete = False
chart.y_axis.delete = False

chart.x_axis.tickLblPos = "low"
chart.y_axis.tickLblPos = "nextTo"

# X軸の表示間隔：2日
chart.x_axis.majorUnit = 2

# X軸の表示範囲を実データに合わせる
first_date = df_chart.index.min()

chart.x_axis.scaling.min = to_excel(
    first_date.to_pydatetime()
)

# X軸の日付表示
chart.x_axis.numFmt = "m/d"

# 折れ線＋プロット点
chart.scatterStyle = "lineMarker"

xvalues = Reference(
    ws_data,
    min_col=1,
    min_row=2,
    max_row=1 + len(df_chart)
)

for col in range(
    2,
    2 + len(df_chart.columns)
):
    values = Reference(
        ws_data,
        min_col=col,
        min_row=1,
        max_row=1 + len(df_chart)
    )

    series = Series(
        values,
        xvalues=xvalues,
        title_from_data=True
    )

    series.smooth = False

    series.marker = Marker(
        symbol="circle",
        size=5
    )

    chart.series.append(series)

# グリッド線
chart.x_axis.majorGridlines = ChartLines(
    spPr=GraphicalProperties(
        ln=LineProperties(
            #solidFill="D9D9D9",
            prstDash="dot"
        )
    )
)


chart.y_axis.majorGridlines = ChartLines(
    spPr=GraphicalProperties(
        ln=LineProperties(
            solidFill="D9D9D9",
            #prstDash="dot"
        )
    )
)

chart.legend.position = "r"

chart.layout = Layout(
    manualLayout=ManualLayout(
        x=0,
        y=0,
        w=0.7,
        h=0.8
    )
)

chart.width = 18
chart.height = 10

ws_report.add_chart(
    chart,
    "E2"
)

# グラフ用データは通常利用者に見せない
ws_data.sheet_state = "hidden"


# ============================================================
# 保存
# ============================================================

try:
    wb.save(output_path)

except PermissionError:
    stop_with_errors([
        f"{output_path.name}を保存できませんでした。"
        "Excelで同じファイルを開いている場合は閉じてください。"
    ])

except Exception as e:
    stop_with_errors([
        f"レポート保存中にエラーが発生しました：{e}"
    ])


print()
print("処理が完了しました。")
print("製品別売上レポートを output フォルダに保存しました。")
print()

input("Enterキーを押すと終了します。")
